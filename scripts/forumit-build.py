"""Build the Word version of the forum доклад from its markdown source.

    python scripts/forumit-build.py

The markdown file in `docs/forumit-2026/` is the source; this script is the only
writer of `docs/forumit-2026/Доклад.docx`. Formatting follows the forum's own
sample (`doc/doclad-it-a-forum.doc`): A4 portrait, Times New Roman 14 pt, single
line spacing, margins 20 mm left and right, 20 mm top, 32 mm bottom, 10 mm
first-line indent, automatic hyphenation, justified body. The heading block of
an article is УДК, names, affiliation with the city and the country, contact;
there is no title page, no table of contents and no annotation. Table captions
and table cells are set to 12 pt the way the sample asks.

Two rules the script enforces instead of trusting the writer:

* no em dash anywhere in the document text (the workspace bans it);
* every heading-block key is present, so a filled article cannot silently lose
  its author or its affiliation.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

HERE = Path(__file__).resolve().parent.parent
SOURCE = HERE / "docs" / "forumit-2026" / "Доклад.md"
TARGET = HERE / "docs" / "forumit-2026" / "Доклад.docx"

FONT = "Times New Roman"
BODY_SIZE = Pt(14)
SMALL_SIZE = Pt(12)
# The forum's sample: 2 cm left and right, 2 cm top, 3,2 cm bottom.
LEFT_MARGIN = Cm(2)
RIGHT_MARGIN = Cm(2)
TOP_MARGIN = Cm(2)
BOTTOM_MARGIN = Cm(3.2)
FIRST_LINE = Cm(1)
# Text width of an A4 page under these margins: 210 - 20 - 20 = 170 mm, which
# is 9638 twips. Word lays a plain table out 108 twips left of the text
# margin, so the indent below cancels that offset. Both numbers were measured
# on the exported pages (frame 56,7..538,6 pt) and then corrected twice: the
# plain indent left the frame 2,3 pt right of the margin, and a 9639-twip wide
# table ended 2,9 pt past it.
TABLE_INDENT_TWIPS = 62
TABLE_WIDTH_TWIPS = 9627
# No figure may be placed taller than this, so that it still fits on one page.
FIGURE_MAX_HEIGHT = Cm(17)

KEYS = ("УДК:", "АВТОРЫ:", "ОРГАНИЗАЦИЯ:", "МЕСТО:", "КОНТАКТ:")
CAPTION = re.compile(r"^Таблица\s+\d+\.")
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


# --------------------------------------------------------------------------- #
# document skeleton
# --------------------------------------------------------------------------- #

def base_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = TOP_MARGIN
    section.bottom_margin = BOTTOM_MARGIN
    section.left_margin = LEFT_MARGIN
    section.right_margin = RIGHT_MARGIN

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = BODY_SIZE
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    fmt = normal.paragraph_format
    fmt.line_spacing = 1.0
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.first_line_indent = FIRST_LINE
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name in ("Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = FONT
        style.font.size = BODY_SIZE
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        pf = style.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.space_before = Pt(10)
        pf.space_after = Pt(6)
        pf.keep_with_next = True
        pf.line_spacing = 1.0

    # The forum sample asks for automatic hyphenation.
    settings = doc.settings.element
    for existing in settings.findall(qn("w:autoHyphenation")):
        settings.remove(existing)
    hyphenation = OxmlElement("w:autoHyphenation")
    hyphenation.set(qn("w:val"), "true")
    settings.append(hyphenation)
    return doc


def page_number_footer(section) -> None:
    """PAGE field centred at the bottom of the page."""
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run()
    for element, attrs, text in (
        ("w:fldChar", {"w:fldCharType": "begin"}, None),
        ("w:instrText", {"xml:space": "preserve"}, " PAGE "),
        ("w:fldChar", {"w:fldCharType": "end"}, None),
    ):
        node = OxmlElement(element)
        for key, value in attrs.items():
            node.set(qn(key), value)
        if text is not None:
            node.text = text
        run._r.append(node)
    run.font.name = FONT
    run.font.size = SMALL_SIZE


# --------------------------------------------------------------------------- #
# inline and block rendering
# --------------------------------------------------------------------------- #

def plain(paragraph, text: str) -> None:
    """Add `text` to `paragraph`, honouring **bold** and `code` spans."""
    for chunk in INLINE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            paragraph.add_run(chunk[2:-2]).bold = True
        elif chunk.startswith("`") and chunk.endswith("`"):
            run = paragraph.add_run(chunk[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(12)
        else:
            paragraph.add_run(chunk)


def body_paragraph(doc, text: str):
    paragraph = doc.add_paragraph()
    plain(paragraph, text)
    return paragraph


def centered(doc, text: str, *, bold: bool = False, italic: bool = False,
             size: Pt | None = None, indent: bool = False):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = FIRST_LINE if indent else Cm(0)
    plain(paragraph, text)
    for run in paragraph.runs:
        run.bold = run.bold or bold
        run.italic = run.italic or italic
        if size is not None:
            run.font.size = size
            run.font.name = FONT
    return paragraph


def caption(doc, text: str):
    """A table caption: 12 pt, centred, on its own line above the table."""
    return centered(doc, text, bold=False, size=SMALL_SIZE)


def align_table(table) -> None:
    """Keep a table inside the same 2 cm margins as the text."""
    tbl_pr = table._tbl.tblPr
    for existing in tbl_pr.findall(qn("w:tblInd")):
        tbl_pr.remove(existing)
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), str(TABLE_INDENT_TWIPS))
    indent.set(qn("w:type"), "dxa")
    marker = tbl_pr.find(qn("w:tblLook"))
    if marker is not None:
        marker.addprevious(indent)
    else:
        tbl_pr.append(indent)

    for existing in tbl_pr.findall(qn("w:tblCellMar")):
        tbl_pr.remove(existing)
    cell_margin = OxmlElement("w:tblCellMar")
    for side in ("left", "right"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), "56")
        node.set(qn("w:type"), "dxa")
        cell_margin.append(node)
    tbl_pr.append(cell_margin)

    for existing in tbl_pr.findall(qn("w:tblW")):
        tbl_pr.remove(existing)
    width = OxmlElement("w:tblW")
    width.set(qn("w:w"), str(TABLE_WIDTH_TWIPS))
    width.set(qn("w:type"), "dxa")
    tbl_pr.append(width)


def add_table(doc, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.autofit = True
    align_table(table)
    for r, cells in enumerate(rows):
        for c, value in enumerate(cells):
            paragraph = table.cell(r, c).paragraphs[0]
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            plain(paragraph, value)
            for run in paragraph.runs:
                run.font.size = SMALL_SIZE
                run.font.name = FONT
                if r == 0:
                    run.bold = True
    doc.add_paragraph()


def figure_width(path: Path) -> Cm:
    """Width for a figure, capped so its height still fits on one page."""
    try:
        from PIL import Image
    except ImportError:
        return Cm(16)
    with Image.open(path) as image:
        ratio = image.height / image.width
    if ratio <= 0:
        return Cm(16)
    allowed = Cm(FIGURE_MAX_HEIGHT.cm / ratio)
    return allowed if allowed < Cm(16) else Cm(16)


def add_figure(doc, path: Path) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=figure_width(path))


def is_table_row(line: str) -> bool:
    return line.startswith("|") and line.endswith("|")


def split_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip("|").split("|")]


# --------------------------------------------------------------------------- #
# main
# --------------------------------------------------------------------------- #

def read_source() -> tuple[dict[str, str], list[str]]:
    """Heading-block fields and the body lines, comments removed."""
    text = SOURCE.read_text(encoding="utf-8")
    text = re.sub(r"<!--.*?-->", "", text, flags=re.S)
    lines = [line.rstrip() for line in text.splitlines()]

    title = ""
    fields: dict[str, str] = {}
    body_start = 0
    for index, line in enumerate(lines):
        if line.startswith("# ") and not title:
            title = line[2:].strip()
            continue
        if any(line.startswith(key) for key in KEYS):
            key, _, value = line.partition(":")
            fields[key.strip() + ":"] = value.strip()
            continue
        if line.startswith("## ") and title:
            body_start = index
            break
    return {"title": title, **fields}, lines[body_start:]


def build() -> int:
    head, body = read_source()
    missing = [key for key in KEYS if key not in head]
    if missing:
        print(f"FAIL: heading block is missing {missing}", file=sys.stderr)
        return 1

    doc = base_document()
    page_number_footer(doc.sections[0])

    # ---- heading block --------------------------------------------------- #
    udk = doc.add_paragraph()
    udk.paragraph_format.first_line_indent = Cm(0)
    udk.alignment = WD_ALIGN_PARAGRAPH.LEFT
    plain(udk, head["УДК:"])

    centered(doc, head["title"], bold=True)
    centered(doc, head["АВТОРЫ:"])
    centered(doc, f"{head['ОРГАНИЗАЦИЯ:']}, {head['МЕСТО:']}", italic=True)
    centered(doc, head["КОНТАКТ:"], size=SMALL_SIZE)
    doc.add_paragraph()

    # ---- body ------------------------------------------------------------ #
    figures = 0
    tables = 0
    index = 0
    while index < len(body):
        line = body[index]

        if not line.strip():
            index += 1
            continue

        image = re.match(r"!\[(.*?)\]\((.*?)\)", line)
        if image:
            path = SOURCE.parent / image.group(2)
            if not path.exists():
                print(f"FAIL: missing figure {path}", file=sys.stderr)
                return 1
            add_figure(doc, path)
            if image.group(1):
                caption(doc, image.group(1))
            figures += 1
            index += 1
            continue

        if is_table_row(line) and index + 1 < len(body) and set(body[index + 1]) <= set("|-: "):
            rows: list[list[str]] = []
            cursor = index
            while cursor < len(body) and is_table_row(body[cursor]):
                if not set(body[cursor]) <= set("|-: "):
                    rows.append(split_row(body[cursor]))
                cursor += 1
            width = max(len(row) for row in rows)
            rows = [row + [""] * (width - len(row)) for row in rows]
            add_table(doc, rows)
            tables += 1
            index = cursor
            continue

        if CAPTION.match(line):
            paragraph = caption(doc, line.strip())
            paragraph.paragraph_format.space_before = Pt(6)
            index += 1
            continue

        if line.startswith("### "):
            plain(doc.add_heading(level=2), line[4:].strip())
            index += 1
            continue

        if line.startswith("## "):
            plain(doc.add_heading(level=1), line[3:].strip())
            index += 1
            continue

        if re.match(r"^\d+\.\s", line):
            paragraph = body_paragraph(doc, line.strip())
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.left_indent = Cm(1)
            index += 1
            continue

        if line.startswith("- "):
            paragraph = body_paragraph(doc, "• " + line[2:].strip())
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.left_indent = Cm(1)
            index += 1
            continue

        body_paragraph(doc, line.strip())
        index += 1

    # ---- the two gates --------------------------------------------------- #
    problems: list[str] = []
    for paragraph in doc.paragraphs:
        if "\u2014" in paragraph.text:
            problems.append(f"em dash in: {paragraph.text[:60]}")
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if "\u2014" in cell.text:
                    problems.append(f"em dash in table: {cell.text[:40]}")
    if problems:
        for problem in problems:
            print("FAIL:", problem, file=sys.stderr)
        return 1

    doc.save(TARGET)
    print(f"saved: {TARGET}")
    print(f"paragraphs: {len(doc.paragraphs)}  tables: {tables}  figures: {figures}")
    print(f"authors: {head['АВТОРЫ:']}")
    if "[" in head["КОНТАКТ:"]:
        print("note: the contact line still holds a placeholder to fill in by hand")
    return 0


if __name__ == "__main__":
    raise SystemExit(build())
